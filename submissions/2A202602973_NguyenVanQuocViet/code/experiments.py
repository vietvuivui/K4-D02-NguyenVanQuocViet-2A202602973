"""experiments.py - điều phối Bước 1-5: chạy (hoặc nạp lại) thí nghiệm, quy tắc chọn trên VAL,
suy luận, chung kết trên test, xuất results.xlsx và hình cho báo cáo.

Mọi lựa chọn (backbone, công thức, phương pháp suy luận) chỉ dùng số liệu trên VAL theo các quy tắc
ghi ở đầu file, và in ra lý do. Test chỉ được chạm ở `final_predict` (Bước 4), mỗi seed đúng một lần.

Tiết kiệm GPU: `run_or_load` bỏ qua lần chạy đã xong (Colab bị ngắt thì chạy lại ô là tiếp tục), và
dùng lại một lần chạy cũ có CÙNG cấu hình (chỉ khác exp_id), ví dụ T00 = backbone đã chọn ở Bước 1,
F01 seed 0 = công thức tốt nhất ở Bước 2. Lần dùng lại được ghi rõ trong summary ("reused_from").
"""
from __future__ import annotations

import copy
import dataclasses
import json
import shutil
from pathlib import Path

import numpy as np
import pandas as pd
import torch

import benchmark as Bm
import dataset as D
import inference as I
import model as M
from train import Config, curve_path, plot_curves, pred_path, run, run_dir, softmax_np  # thêm gốc repo vào sys.path
from eval import compute_metrics, mean_std, save_predictions  # noqa: E402

# ----------------------------------------------------------------------------- quy tắc chọn trên VAL
NOISE_F1 = 0.005          # Bước 1 (chưa đo nhiễu): chênh lệch macro-F1 val nhỏ hơn mức này coi là "không phân biệt được"
INFER_MIN_GAIN = 0.002    # phương pháp suy luận phải hơn I00 ít nhất mức này mới được chọn cho chung kết
LATENCY_BUDGET_MS = 100.0  # ngân sách thời gian thực, p95 batch 1 (slide trang 61)
NOISE_FLOOR = 0.003       # sàn ngưỡng nhiễu ở Bước 2 (slide trang 59: dưới ~0,3 điểm là nhiễu)
NOISE_SEEDS = (0, 1, 2)   # seed của T00 dùng để đo nhiễu (cũng là mốc T00 ở Bước 4)
RARE = {"Chinee Apple": 0, "Snake Weed": 7}


def _is_rare(names: pd.Series) -> pd.Series:
    """Tên lớp có thể là "Chinee Apple" (dataset.py) hoặc "Chinee apple" (labels.csv/eval.py): so không phân biệt hoa thường."""
    return names.astype(str).str.lower().isin([k.lower() for k in RARE])

# Ghim tag trọng số (timm "kiến_trúc.tag"): tag mặc định đổi theo phiên bản timm, và mặc định của convnext_tiny
# là in12k_ft_in1k (tiền huấn luyện ImageNet-12k) -> không công bằng. Tất cả dưới đây chỉ dùng ImageNet-1k.
BACKBONES = [
    ("B01", "resnet50.a1_in1k"),                   # ResNet (mốc)
    ("B02", "convnext_tiny.fb_in1k"),              # ConvNeXt
    ("B03", "deit_small_patch16_224.fb_in1k"),     # transformer
    ("B04", "efficientnet_b0.ra_in1k"),            # nhẹ
    ("B05", "mobilenetv3_large_100.ra_in1k"),      # nhẹ
]

# (exp_id, trục, mô tả "khác T00 ở điểm nào", ghi đè Config). Mỗi dòng khác T00 đúng MỘT yếu tố.
ABLATIONS = [
    ("T01", "A", "init=frozen (chỉ train head)", dict(init="frozen")),
    ("T02", "B", "aug=trivial (TrivialAugmentWide)", dict(aug="trivial")),
    ("T03", "B", "aug=flip_rot (lật dọc + xoay 90°)", dict(aug="flip_rot")),
    ("T04", "B", "CutMix alpha=1.0", dict(mix="cutmix", mix_alpha=1.0)),
    ("T05", "C", "loss=label smoothing 0.1", dict(loss="ls", label_smoothing=0.1)),
    ("T06", "C", "loss=focal gamma=2", dict(loss="focal", focal_gamma=2.0)),
    ("T07", "C", "loss=CE trọng số 1/n_c (train)", dict(loss="ce_weighted", class_weight_beta=0.0)),
    ("T08", "F", "EMA decay=0.995", dict(ema_decay=0.995)),  # ~2000 bước ở 12 epoch: 0.995^2000 ≈ 0,004% trọng số khởi tạo còn lại
    ("T09", "D", "sampler cân bằng lớp (oversample)", dict(sampler="balanced")),
]
COMBO_ID = "T10"

# các trường không ảnh hưởng kết quả huấn luyện: bỏ qua khi so "cùng cấu hình"
_NON_RECIPE = {"exp_id", "desc", "save_test_predictions", "num_workers", "save_checkpoint", "images_dir",
               "labels_dir", "out_dir", "pred_dir", "curves_dir", "cache_images", "cache_dir"}


def _slug(text: str) -> str:
    """Tên file ASCII: bỏ dấu tiếng Việt, ký tự lạ thành '_'."""
    import unicodedata
    text = unicodedata.normalize("NFKD", text.replace("đ", "d").replace("Đ", "D"))
    text = "".join(c for c in text if not unicodedata.combining(c))
    keep = "".join(c if (c.isascii() and c.isalnum()) or c in "._-" else "_" for c in text)
    return "_".join(p for p in keep.split("_") if p)[:40]


def _read_json(p: Path) -> dict:
    return json.loads(Path(p).read_text(encoding="utf-8"))


def _write_json(p: Path, d: dict) -> None:
    Path(p).write_text(json.dumps(d, indent=2, ensure_ascii=False, default=float), encoding="utf-8")


# ============================================================================= chạy / nạp lại
def _recipe(d: dict) -> dict:
    return {k: v for k, v in d.items() if k not in _NON_RECIPE}


def find_equivalent(cfg: Config) -> Path | None:
    """Tìm lần chạy đã xong có cùng công thức + seed (khác exp_id)."""
    want = _recipe(dataclasses.asdict(cfg))
    for cj in sorted(Path(cfg.out_dir).glob("*/seed*/config.json")):
        if not (cj.parent / "summary.json").exists():
            continue
        d = _read_json(cj)
        if d.get("exp_id") != cfg.exp_id and _recipe(d) == want:
            return cj.parent
    return None


def _alias_run(src: Path, cfg: Config) -> dict:
    """Dùng lại lần chạy `src` dưới exp_id mới: chép log/logit (không chép checkpoint), vẽ lại đường cong."""
    dst = run_dir(cfg)
    shutil.copytree(src, dst, dirs_exist_ok=True, ignore=shutil.ignore_patterns("best*.pt"))
    old = _read_json(src / "summary.json")
    s = {**old, "exp_id": cfg.exp_id, "run_dir": str(dst), "reused_from": f"{old['exp_id']} seed{old['seed']}"}
    _write_json(dst / "config.json", dataclasses.asdict(cfg))
    z = np.load(dst / "val_logits.npz")
    save_predictions(pred_path(cfg, "val"), z["filenames"], z["y_true"], softmax_np(z["logits"]))
    hist = pd.read_csv(dst / "history.csv").to_dict("records")
    lrs = np.load(dst / "lr_steps.npy").tolist() if (dst / "lr_steps.npy").exists() else None
    title = (f"{cfg.exp_id} · {cfg.backbone} · seed {cfg.seed} · best epoch {s['best_epoch']} "
             f"(val macro-F1 {s['val_macro_f1']:.4f}) · cùng lần chạy với {s['reused_from']}")
    plot_curves(hist, curve_path(cfg), title, lrs)
    _write_json(dst / "summary.json", s)
    print(f"[{cfg.exp_id} seed{cfg.seed}] dùng lại {s['reused_from']} (cùng công thức và seed)")
    return s


def run_or_load(cfg: Config, reuse: bool = True) -> dict:
    """Đã chạy xong -> đọc summary; có lần chạy cùng cấu hình -> dùng lại; còn lại -> train.run(cfg)."""
    summ = run_dir(cfg) / "summary.json"
    if summ.exists():
        saved = _recipe(_read_json(run_dir(cfg) / "config.json"))
        want = _recipe(dataclasses.asdict(cfg))
        if saved != want:  # cùng exp_id/seed nhưng cấu hình khác -> không được nạp nhầm kết quả cũ
            diff = {k: (saved.get(k), want.get(k)) for k in set(saved) | set(want) if saved.get(k) != want.get(k)}
            raise ValueError(f"{run_dir(cfg)} đã có kết quả với cấu hình KHÁC (đã lưu, yêu cầu): {diff}. "
                             "Xoá thư mục đó hoặc đổi exp_id.")
        return _read_json(summ)
    src = find_equivalent(cfg) if reuse else None
    return _alias_run(src, cfg) if src is not None else run(cfg)


def load_config(summary: dict) -> Config:
    return Config(**_read_json(Path(summary["run_dir"]) / "config.json"))


def load_trained(summary: dict, device) -> torch.nn.Module:
    """Dựng lại kiến trúc (không tải trọng số ImageNet) rồi nạp checkpoint tốt nhất của lần chạy."""
    cfg = load_config(summary)
    m = M.build_model(cfg.backbone, pretrained=False, num_classes=D.NUM_CLASSES, drop_rate=cfg.drop_rate,
                      init="scratch")
    ckpt = summary.get("checkpoint") or str(Path(summary["run_dir"]) / "best.pt")
    m.load_state_dict(torch.load(ckpt, map_location="cpu"))
    return m.to(device).eval()


def _norm_of(model) -> tuple:
    pc = getattr(model, "pretrained_cfg", {}) or {}
    return pc.get("mean", D.IMAGENET_MEAN), pc.get("std", D.IMAGENET_STD)


def _val_logits(summary: dict):
    z = np.load(Path(summary["run_dir"]) / "val_logits.npz")
    return z["filenames"], z["y_true"], z["logits"]


def _metrics(y, probs) -> dict:
    return compute_metrics(np.asarray(y), probs.argmax(1), probs)


def _empty_cache() -> None:
    if torch.cuda.is_available():
        torch.cuda.empty_cache()


# ============================================================================= Bước 1
def step1_backbones(paths: dict, device, backbones=BACKBONES, **overrides) -> pd.DataFrame:
    """Huấn luyện mỗi backbone bằng công thức nền T00 (cùng seed 0) + đo độ trễ sơ bộ batch 1 FP32."""
    rows = []
    for exp_id, bb in backbones:
        cfg = Config(exp_id=exp_id, backbone=bb, desc=bb, seed=0, **overrides, **paths)
        s = run_or_load(cfg)
        if "latency_b1_p50_ms" not in s:
            m = load_trained(s, device)
            lat = Bm.latency_report(m, 1, cfg.img_size, "fp32", device.type, warmup=10, iters=100)
            s.update(latency_b1_p50_ms=lat["p50"], latency_b1_p95_ms=lat["p95"], latency_gpu=lat["gpu"])
            _write_json(Path(s["run_dir"]) / "summary.json", s)
            del m
            _empty_cache()
        rows.append({**s, "img_size": cfg.img_size, "epochs": cfg.epochs})
    df = pd.DataFrame(rows)
    df["note"] = ("công thức nền T00, seed 0, 1 seed; độ trễ sơ bộ FP32 batch 1 (warmup 10, 100 lần); "
                  "GMAC đếm bằng torch FlopCounterMode (FLOPs/2)")
    return df[["exp_id", "backbone", "weights_tag", "params_M", "gmacs", "img_size", "epochs", "seed",
               "best_epoch", "val_macro_f1", "val_top1", "val_balanced_acc", "train_time_per_epoch_s",
               "latency_b1_p50_ms", "latency_b1_p95_ms", "latency_gpu", "note"]]


def backbone_analysis(bb_df: pd.DataFrame, paths: dict, overfit_tol: float = 0.05) -> tuple[pd.DataFrame, dict]:
    """Thêm cột phân tích từ history.csv của mỗi B0x và tính tương quan FLOPs với thời gian/độ trễ.

    - epochs_to_99pct: epoch đầu tiên đạt >= 99% macro-F1 val tốt nhất (hội tụ nhanh hay chậm)
    - val_loss_min_epoch, val_loss_rise: val loss ở epoch cuối trừ val loss nhỏ nhất
    - overfit: val loss tăng > overfit_tol sau điểm thấp nhất trong khi train loss vẫn giảm
    Tương quan Spearman (thứ hạng, 5 điểm nên chỉ mang tính tham khảo).
    """
    df = bb_df.copy()
    extra = []
    for r in df.itertuples():
        h = pd.read_csv(Path(paths["out_dir"]) / r.exp_id / "seed0" / "history.csv")
        best = h["val_macro_f1"].max()
        i_min = int(h["val_loss"].idxmin())
        rise = float(h["val_loss"].iloc[-1] - h["val_loss"].iloc[i_min])
        train_falls = bool(h["train_loss"].iloc[-1] < h["train_loss"].iloc[i_min])
        extra.append({"epochs_to_99pct": int(h.loc[h["val_macro_f1"] >= 0.99 * best, "epoch"].iloc[0]),
                      "val_loss_min_epoch": int(h["epoch"].iloc[i_min]), "val_loss_rise": rise,
                      "overfit": bool(rise > overfit_tol and train_falls),
                      "train_val_gap_last": float(h["val_loss"].iloc[-1] - h["train_loss"].iloc[-1])})
    df = pd.concat([df.reset_index(drop=True), pd.DataFrame(extra)], axis=1)
    corr = {f"spearman(GMAC, {c})": float(df["gmacs"].corr(df[c], method="spearman"))
            for c in ("latency_b1_p50_ms", "train_time_per_epoch_s")}
    corr["spearman(params, latency_b1_p50_ms)"] = float(df["params_M"].corr(df["latency_b1_p50_ms"], method="spearman"))
    return df, corr


def plot_backbone_curves(bb_df: pd.DataFrame, paths: dict, path: str | Path) -> None:
    """Chồng macro-F1 val và val loss theo epoch của mọi backbone: so tốc độ hội tụ và quá khớp."""
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(1, 2, figsize=(12, 4))
    for r in bb_df.itertuples():
        h = pd.read_csv(Path(paths["out_dir"]) / r.exp_id / "seed0" / "history.csv")
        ax[0].plot(h["epoch"], h["val_macro_f1"], "o-", ms=3, label=f"{r.exp_id} {r.backbone}")
        ax[1].plot(h["epoch"], h["val_loss"], "o-", ms=3, label=f"{r.exp_id}")
    ax[0].set(xlabel="epoch", ylabel="macro-F1 val", title="Hội tụ: macro-F1 val theo epoch")
    ax[1].set(xlabel="epoch", ylabel="val loss (CE)", title="Quá khớp: val loss theo epoch")
    for a in ax:
        a.grid(alpha=0.3)
        a.legend(fontsize=7)
    fig.tight_layout()
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=130)
    plt.close(fig)


def select_backbone(bb_df: pd.DataFrame, override: str | None = None) -> tuple[str, str]:
    """Quy tắc: lấy các backbone có macro-F1 val trong khoảng NOISE_F1 của mức cao nhất (không phân biệt
    được với 1 seed), rồi chọn cái có độ trễ batch 1 thấp nhất. Trả về (backbone, lý do)."""
    if override:
        return override, f"chọn thủ công: {override}"
    best = bb_df["val_macro_f1"].max()
    cands = bb_df[bb_df["val_macro_f1"] >= best - NOISE_F1].sort_values("latency_b1_p50_ms")
    pick = cands.iloc[0]
    top = bb_df.loc[bb_df["val_macro_f1"].idxmax()]
    reason = (f"macro-F1 val cao nhất: {top.backbone} {top.val_macro_f1:.4f}. Các backbone trong khoảng "
              f"{NOISE_F1} (1 seed, không phân biệt được): {', '.join(cands.backbone)}. "
              f"Chọn {pick.backbone} (macro-F1 {pick.val_macro_f1:.4f}, p50 batch-1 {pick.latency_b1_p50_ms:.1f} ms, "
              f"{pick.params_M:.1f}M tham số) vì nhanh nhất trong nhóm đó.")
    return pick.backbone, reason


# ============================================================================= Bước 2
def _rare_f1(y, logits) -> dict:
    m = _metrics(y, softmax_np(logits))
    return {f"f1_{k.split()[0].lower()}": float(m["f1"][i]) for k, i in RARE.items()}


def _training_row(s: dict, exp_id: str, axis: str, change: str, base_f1: float | None) -> dict:
    _, y, lg = _val_logits(s)
    return {"exp_id": exp_id, "backbone": s["backbone"], "axis": axis, "change_vs_T00": change, "seed": s["seed"],
            "val_macro_f1": s["val_macro_f1"], "val_top1": s["val_top1"],
            "delta_vs_T00": s["val_macro_f1"] - base_f1 if base_f1 is not None else 0.0,
            **_rare_f1(y, lg), "best_epoch": s["best_epoch"], "reused_from": s.get("reused_from", "")}


def _t00(paths: dict, backbone: str, seed: int, **overrides) -> dict:
    return run_or_load(Config(exp_id="T00", backbone=backbone, desc=f"{_slug(backbone)}_baseline", seed=seed,
                              **overrides, **paths))


def measure_noise(paths: dict, backbone: str, seeds=None, **overrides) -> dict:
    """Đo nhiễu thật: T00 với nhiều seed (seed 0 dùng lại B0x; các seed khác cũng chính là mốc T00 của
    Bước 4, nên tổng số lần train không đổi). Ngưỡng = max(std mẫu ddof=1, NOISE_FLOOR)."""
    seeds = tuple(seeds or NOISE_SEEDS)
    runs = [_t00(paths, backbone, sd, **overrides) for sd in seeds]
    f1 = np.array([r["val_macro_f1"] for r in runs])
    std = float(f1.std(ddof=1)) if len(f1) > 1 else float("nan")
    thr = max(std, NOISE_FLOOR) if np.isfinite(std) else NOISE_F1
    return {"seeds": list(seeds), "val_macro_f1": f1.tolist(), "mean": float(f1.mean()), "std": std,
            "threshold": thr, "runs": runs}


def _vs_noise(delta: float, thr: float) -> str:
    if delta > thr:
        return "vượt nhiễu (tốt hơn)"
    if delta < -thr:
        return "kém hơn rõ"
    return "không phân biệt được"


def annotate_noise(tr_df: pd.DataFrame, noise: dict | None) -> pd.DataFrame:
    """Thêm cột noise_threshold và vs_noise (kết luận của từng dòng so với nhiễu)."""
    thr = noise["threshold"] if noise else NOISE_F1
    df = tr_df.copy()
    df["noise_threshold"] = thr
    df["vs_noise"] = [("mốc" if (r.exp_id == "T00" and r.seed == 0) else
                       "đo nhiễu" if r.exp_id == "T00" else _vs_noise(r.delta_vs_T00, thr))
                      for r in df.itertuples()]
    df["note"] = [(f"dùng lại lần chạy {r.reused_from} (cùng cấu hình, cùng seed)" if r.reused_from else "")
                  + ("; 1 seed" if r.exp_id != "T00" else "") for r in df.itertuples()]
    return df


def step2_training(paths: dict, backbone: str, ablations=ABLATIONS, noise: dict | None = None,
                   **overrides) -> pd.DataFrame:
    """T00 (nền, seed 0) + mỗi ablation khác T00 đúng một yếu tố, cùng backbone, seed 0.
    Nếu có `noise` (measure_noise), thêm dòng T00 các seed khác và cột so Δ với ngưỡng nhiễu."""
    s0 = _t00(paths, backbone, 0, **overrides)
    rows = [_training_row(s0, "T00", "-", "công thức nền", None)]
    for r in (noise or {}).get("runs", []):
        if r["seed"] != 0:
            rows.append(_training_row(r, "T00", "-", f"công thức nền, seed {r['seed']} (đo nhiễu)", s0["val_macro_f1"]))
    for exp_id, axis, change, ov in ablations:
        s = run_or_load(Config(exp_id=exp_id, backbone=backbone, desc=_slug(change), seed=0,
                               **{**overrides, **ov}, **paths))
        rows.append(_training_row(s, exp_id, axis, change, s0["val_macro_f1"]))
    return annotate_noise(pd.DataFrame(rows), noise)


def select_combination(tr_df: pd.DataFrame, ablations=ABLATIONS) -> tuple[dict, str]:
    """Quy tắc kết hợp. Mọi T0x so với CÙNG T00 (không đổi nền giữa chừng), nên thứ tự các trục không ảnh hưởng.
    Mỗi trục lấy biến thể có Δ lớn nhất; ghép các trục mà biến thể đó THẮNG RÕ (Δ > ngưỡng nhiễu đo được).
    Nếu ít hơn 2 trục thắng rõ, nới thành Δ > 0; nếu vẫn ít hơn 2, ghép 2 trục có Δ lớn nhất, để vẫn có
    một thí nghiệm kiểm tra tính cộng dồn (GUIDE mục 3.1, ý 4). Trả về (ghi đè Config, lý do)."""
    ov_of = {e: ov for e, _, _, ov in ablations}
    thr = float(tr_df["noise_threshold"].iloc[0]) if "noise_threshold" in tr_df else NOISE_F1
    abl = tr_df[tr_df.exp_id.isin(ov_of)].sort_values("delta_vs_T00", ascending=False)
    per_axis = abl.groupby("axis", sort=False).head(1)
    rule = f"Δ > ngưỡng nhiễu {thr:.4f}"
    chosen = per_axis[per_axis.delta_vs_T00 > thr]
    if len(chosen) < 2:
        chosen, rule = per_axis[per_axis.delta_vs_T00 > 0], f"ít hơn 2 trục vượt nhiễu {thr:.4f} -> nới thành Δ > 0"
    if len(chosen) < 2:
        chosen, rule = per_axis.head(2), "ít hơn 2 trục có Δ > 0 -> ghép 2 trục có Δ lớn nhất để kiểm tra cộng dồn"
    combo = {}
    for e in chosen.exp_id:
        combo.update(ov_of[e])
    reason = (f"Mọi T0x so với cùng T00 (không đổi nền giữa chừng). Quy tắc chọn: {rule}. Ghép "
              + " + ".join(f"{r.exp_id} ({r.change_vs_T00}, Δ={r.delta_vs_T00:+.4f})" for r in chosen.itertuples())
              + f" thành {COMBO_ID}.")
    return combo, reason


def step2_combo(paths: dict, backbone: str, combo: dict, **overrides) -> dict:
    return run_or_load(Config(exp_id=COMBO_ID, backbone=backbone, desc="combo", seed=0,
                              **{**overrides, **combo}, **paths))


def select_recipe(tr_df: pd.DataFrame, ablations=ABLATIONS, combo: dict | None = None) -> tuple[str, dict, str]:
    """Công thức chung kết = dòng seed 0 có macro-F1 val cao nhất trong T00..T10. Trả về (exp_id, ghi đè, lý do)."""
    ov_of = {"T00": {}, **{e: ov for e, _, _, ov in ablations}, COMBO_ID: combo or {}}
    thr = float(tr_df["noise_threshold"].iloc[0]) if "noise_threshold" in tr_df else NOISE_F1
    cand = tr_df[tr_df.seed == 0]
    best = cand.loc[cand["val_macro_f1"].idxmax()]
    base = cand.loc[cand.exp_id == "T00", "val_macro_f1"].item()
    d = best.val_macro_f1 - base
    verdict = "là chính mốc" if best.exp_id == "T00" else _vs_noise(d, thr)
    reason = (f"{best.exp_id} có macro-F1 val cao nhất {best.val_macro_f1:.4f} (T00 seed 0 {base:.4f}, Δ={d:+.4f}); "
              f"so với ngưỡng nhiễu {thr:.4f}: {verdict}. Lựa chọn dựa trên 1 seed; Bước 4 kiểm chứng bằng 3 seed.")
    return best.exp_id, ov_of[best.exp_id], reason


# ============================================================================= Bước 3
METHODS = {
    "I00": dict(kind="single", img_size=224, k=1, label="1 view (resize 256 + center crop 224)"),
    "I01": dict(kind="hflip", space="prob", img_size=224, k=2, label="TTA lật ngang, gộp xác suất"),
    "I03a": dict(kind="hflip", space="logit", img_size=224, k=2, label="TTA lật ngang, gộp logit"),
    "I02": dict(kind="multicrop", space="prob", img_size=224, k=10, label="TTA 5 crop + lật (K=10), gộp xác suất"),
    "I03b": dict(kind="multicrop", space="logit", img_size=224, k=10, label="TTA 5 crop + lật (K=10), gộp logit"),
    "I04_256": dict(kind="single", img_size=256, k=1, label="độ phân giải kiểm tra 256"),
    "I04_288": dict(kind="single", img_size=288, k=1, label="độ phân giải kiểm tra 288"),
    "I04_320": dict(kind="single", img_size=320, k=1, label="độ phân giải kiểm tra 320"),
}


def method_logits(model, cfg: Config, df: pd.DataFrame, method: dict, device):
    """Logit theo từng view của một phương pháp suy luận. Trả về (filenames, y, [logits_view]*K)."""
    mean, std = _norm_of(model)
    if method["kind"] == "multicrop":   # ảnh 256 nguyên vẹn, cắt 5 crop 224 trên GPU
        tf = D.build_transforms(False, 256, mean=mean, std=std, eval_resize=256)
        views_fn = lambda x: I.views_multicrop(x, method["img_size"], flip=True)  # noqa: E731
    else:
        tf = D.build_transforms(False, method["img_size"], mean=mean, std=std)
        views_fn = I.views_hflip if method["kind"] == "hflip" else (lambda x: [x])  # noqa: E731
    loader = D.make_loader(df, cfg.images_dir, tf, batch_size=64, train=False,
                           num_workers=cfg.num_workers, cache=cfg.cache_dir or False)
    return I.predict_views(model, loader, device, views_fn, amp=cfg.amp)


def combine(view_logits: list, method: dict) -> tuple[np.ndarray, np.ndarray]:
    """Gộp view -> (probs, pseudo-logit dùng cho temperature scaling)."""
    if len(view_logits) == 1:
        lg = view_logits[0]
        return softmax_np(lg), lg
    probs = I.aggregate_views(view_logits, method.get("space", "prob"))
    pseudo = np.mean(view_logits, 0) if method.get("space") == "logit" else np.log(np.clip(probs, 1e-12, None))
    return probs, pseudo


def _ece_crossfit(logits, y, seed: int = 0) -> tuple[float, float]:
    """ECE sau TS ước lượng trung thực hơn: khớp T trên một nửa val, đo ECE trên nửa còn lại (2 chiều)."""
    rng = np.random.default_rng(seed)
    idx = rng.permutation(len(y))
    halves = np.array_split(idx, 2)
    eces, temps = [], []
    for a, b in ((0, 1), (1, 0)):
        T = I.fit_temperature(logits[halves[a]], y[halves[a]])
        p = I.apply_temperature(logits[halves[b]], T)
        eces.append(compute_metrics(y[halves[b]], p.argmax(1), p)["ece"])
        temps.append(T)
    return float(np.mean(eces)), float(np.mean(temps))


def _lat_row(cfg_name, lat, fused: bool) -> dict:
    return {"config": cfg_name, "gpu": lat["gpu"], "dtype": lat["dtype"], "batch": lat["batch"],
            "img_size": lat["img_size"], "k_views": lat["k_views"], "bn_fused": fused,
            "p50_ms": lat["p50"], "p95_ms": lat["p95"], "p99_ms": lat["p99"],
            "images_per_s": lat["images_per_s"], "torch": lat["torch"], "preprocessing_included": False}


# Phân loại chi phí suy luận (để trả lời: TTA/ensemble hợp ngoại tuyến, robot nên dùng thứ không tốn thêm)
def _cost_class(code: str) -> str:
    if code == "I00":
        return "mốc (1 lượt)"
    if code.startswith(("I06", "I07", "I08")):
        return "không tốn thêm"
    if code.startswith("I04"):
        return "1 lượt, FLOPs tăng"
    return "tốn thêm (K lượt / nhiều mô hình)"


def _soup_state(summaries: list[dict]) -> dict:
    """Uniform soup: trung bình trọng số (và buffer BN) của các mô hình cùng kiến trúc, cùng trọng số tiền huấn luyện."""
    states = [torch.load(s.get("checkpoint") or str(Path(s["run_dir"]) / "best.pt"), map_location="cpu")
              for s in summaries]
    out = {}
    for k, v in states[0].items():
        out[k] = (torch.stack([st[k].float() for st in states]).mean(0).to(v.dtype)
                  if v.dtype.is_floating_point else v.clone())
    return out


def step3_inference(recipe_summary: dict, candidates: list[dict], device, cache_csv: str | Path,
                    ema_summary: dict | None = None, seed_runs: list[dict] | None = None,
                    ) -> tuple[pd.DataFrame, pd.DataFrame]:
    """So sánh phương pháp suy luận trên VAL (không train lại), cho model của công thức đã chọn (seed 0).

    candidates : các lần chạy tốt nhất Bước 1-2 để ensemble khác backbone/công thức (I05).
    ema_summary: lần chạy có EMA (T08): so trọng số EMA với trọng số thường CÙNG lần chạy, cùng epoch (I06).
    seed_runs  : T00 seed 0/1/2: ensemble khác seed (I05_seeds) vs model soup (I06_soup), cùng 3 mô hình.
    Độ trễ: batch 1 (p50/p95/p99) và thông lượng batch 32 cho từng phương pháp. Lưu cache CSV.
    """
    cache_csv = Path(cache_csv)
    lat_csv = cache_csv.with_name(cache_csv.stem + "_latency.csv")
    if cache_csv.exists() and lat_csv.exists():
        return pd.read_csv(cache_csv), pd.read_csv(lat_csv)

    cfg = load_config(recipe_summary)
    model = load_trained(recipe_summary, device)
    _, val_df, _ = D.load_split(cfg.labels_dir, cfg.fold)
    who = f"{recipe_summary['exp_id']} seed{recipe_summary['seed']} ({cfg.backbone})"
    rows, lat_rows = [], []
    dev = device.type

    def lat_of(img_size, k, dtype="fp32", m=None, batch=1):
        return Bm.latency_report(m if m is not None else model, batch, img_size, dtype, dev,
                                 warmup=10, iters=100, k_views=k)

    def lat_pair(img_size, k, dtype="fp32", m=None):
        """(độ trễ batch 1, thông lượng batch 32 tính theo ảnh/s)."""
        return lat_of(img_size, k, dtype, m), lat_of(img_size, k, dtype, m, batch=32)["images_per_s"]

    def add(code, method, model_desc, K, img_size, m, lat1, thr32, note="", **extra):
        rows.append({"exp_id": code, "method": method, "model": model_desc, "K": K, "img_size": img_size,
                     "val_macro_f1": m["macro_f1"] if m else np.nan, "val_top1": m["top1"] if m else np.nan,
                     "val_ece": m["ece"] if m else np.nan, "p50_ms": lat1["p50"], "p95_ms": lat1["p95"],
                     "p99_ms": lat1["p99"], "images_per_s_b1": lat1["images_per_s"], "images_per_s_b32": thr32,
                     "note": note, **extra})

    base_lat, base_thr = lat_pair(224, 1)
    for code, meth in METHODS.items():
        try:
            names, y, views = method_logits(model, cfg, val_df, meth, device)
        except Exception as e:  # vd ViT/DeiT không nhận độ phân giải khác 224
            rows.append({"exp_id": code, "method": meth["label"], "model": who, "K": meth["k"],
                         "img_size": meth["img_size"],
                         "note": f"không áp dụng: {type(e).__name__}: {str(e)[:80]}"})
            continue
        probs, _ = combine(views, meth)
        lat1, thr = (base_lat, base_thr) if code == "I00" else lat_pair(meth["img_size"], meth["k"])
        add(code, meth["label"], who, meth["k"], meth["img_size"], _metrics(y, probs), lat1, thr)
        if code == "I00":
            y0, logits0 = y, views[0]

    def ensemble_row(code, label, members, note):
        probs_list, lats, thrs, names = [], [], [], []
        for s in members:
            _, y_, lg_ = _val_logits(s)
            probs_list.append(softmax_np(lg_))
            mm = load_trained(s, device)
            l1, t32 = lat_pair(224, 1, m=mm)
            lats.append(l1)
            thrs.append(t32)
            names.append(f"{s['exp_id']} seed{s['seed']} ({s['backbone']})")
            del mm
            _empty_cache()
        m = _metrics(y_, I.ensemble_probs(probs_list))
        tot = {q: sum(l[q] for l in lats) for q in ("p50", "p95", "p99")}
        tot["images_per_s"] = 1000.0 / tot["p50"]
        add(code, label, " + ".join(names), len(members), 224, m, tot, 1.0 / sum(1.0 / t for t in thrs), note)

    # I05: ensemble khác backbone/công thức (logit val 1-view đã lưu, cùng thứ tự file)
    if len(candidates) >= 2:
        ensemble_row("I05", f"ensemble {len(candidates)} mô hình tốt nhất (trung bình xác suất)", candidates,
                     "độ trễ = tổng độ trễ từng mô hình (chạy tuần tự)")
    # I05_seeds + I06_soup: cùng 3 mô hình T00 khác seed -> ensemble (chi phí xK) vs soup (chi phí x1)
    if seed_runs and len(seed_runs) >= 2:
        ensemble_row("I05_seeds", f"ensemble {len(seed_runs)} seed của T00", seed_runs,
                     "cùng công thức, khác seed; độ trễ = tổng từng mô hình")
        soup_cfg = load_config(seed_runs[0])
        soup = M.build_model(soup_cfg.backbone, pretrained=False, num_classes=D.NUM_CLASSES, init="scratch")
        soup.load_state_dict(_soup_state(seed_runs))
        soup = soup.to(device).eval()
        _, ys, vs = method_logits(soup, soup_cfg, val_df, METHODS["I00"], device)
        add("I06_soup", f"model soup đều {len(seed_runs)} seed của T00 (trung bình trọng số)",
            " + ".join(f"T00 seed{s['seed']}" for s in seed_runs), 1, 224, _metrics(ys, softmax_np(vs[0])),
            *lat_pair(224, 1, m=soup),
            note="chi phí như 1 mô hình; head mỗi seed khởi tạo khác nhau nên soup có thể kém (ghi nhận, không phải lỗi)")
        del soup
        _empty_cache()

    # I06: EMA vs trọng số thường, CÙNG lần chạy T08, cùng epoch tốt nhất
    if ema_summary is not None and ema_summary.get("checkpoint_raw") and Path(ema_summary["checkpoint_raw"]).exists():
        ecfg = load_config(ema_summary)
        for code, ck, label in (("I06", ema_summary["checkpoint"], "trọng số EMA"),
                                ("I06_raw", ema_summary["checkpoint_raw"], "trọng số thường (không EMA)")):
            mm = M.build_model(ecfg.backbone, pretrained=False, num_classes=D.NUM_CLASSES, init="scratch")
            mm.load_state_dict(torch.load(ck, map_location="cpu"))
            mm = mm.to(device).eval()
            _, ye, ve = method_logits(mm, ecfg, val_df, METHODS["I00"], device)
            add(code, f"{label}, {ema_summary['exp_id']} epoch {ema_summary['best_epoch']}",
                f"{ema_summary['exp_id']} seed{ema_summary['seed']}", 1, 224, _metrics(ye, softmax_np(ve[0])),
                base_lat, base_thr, note="cùng lần chạy, cùng epoch; chỉ khác trọng số dùng lúc suy luận")
            del mm
            _empty_cache()

    # I07: temperature scaling trên logit I00 (T khớp trên val)
    T = I.fit_temperature(logits0, y0)
    m_ts = _metrics(y0, I.apply_temperature(logits0, T))
    ece_cf, _ = _ece_crossfit(logits0, y0)
    add("I07", f"temperature scaling (T={T:.3f}, khớp trên val)", who, 1, 224, m_ts, base_lat, base_thr,
        note="ECE sau TS đo trên chính val (in-sample); val_ece_crossfit: khớp T ở nửa val, đo nửa kia",
        val_ece_before=_metrics(y0, softmax_np(logits0))["ece"], val_ece_crossfit=ece_cf, temperature=T)

    # I08: gộp BN + FP16/AMP: độ chính xác và độ trễ
    fused = I.fuse_conv_bn(model, check_input=torch.randn(2, 3, 224, 224, device=device))
    fuse_note = (f"gộp {fused.n_fused} cặp Conv+BN, sai số lớn nhất {fused.fuse_max_abs_diff:.1e}" if fused.n_fused
                 else "kiến trúc không có BatchNorm (LayerNorm): gộp BN không áp dụng, giống hệt I00")
    variants = [("I08_fused_fp32", "gộp BN, FP32", fused, "fp32", True)]
    if dev == "cuda":
        variants += [("I08_amp", "AMP (autocast FP16)", model, "amp", False),
                     ("I08_fp16", "FP16 (model.half())", model, "fp16", False),
                     ("I08_fused_fp16", "gộp BN + FP16", fused, "fp16", True)]
    mean, std = _norm_of(model)
    loader = D.make_loader(val_df, cfg.images_dir, D.build_transforms(False, 224, mean=mean, std=std), batch_size=64,
                           train=False, num_workers=cfg.num_workers, cache=cfg.cache_dir or False)
    for code, label, mm, dtype, is_fused in variants:
        if dtype == "fp16":
            mh = copy.deepcopy(mm).half()
            _, y8, lg8 = I.predict_logits(mh, loader, device, view=lambda t: t.half())
            del mh
        else:
            _, y8, lg8 = I.predict_logits(mm, loader, device, amp=(dtype == "amp"))
        add(code, label, who, 1, 224, _metrics(y8, softmax_np(lg8.astype(np.float64))),
            *lat_pair(224, 1, dtype=dtype, m=mm), note=fuse_note if is_fused else "")

    # Bảng Latency: batch 1 và batch 32, các dtype, có/không gộp BN
    for dtype in (["fp32", "amp", "fp16"] if dev == "cuda" else ["fp32"]):
        for batch in (1, 32):
            lat_rows.append(_lat_row(f"{who} 1-view", lat_of(224, 1, dtype, batch=batch), False))
            lat_rows.append(_lat_row(f"{who} 1-view", lat_of(224, 1, dtype, m=fused, batch=batch), True))
    lat_rows.append(_lat_row(f"{who} TTA lật K=2", lat_of(224, 2), False))
    lat_rows.append(_lat_row(f"{who} TTA 10 crop", lat_of(224, 10), False))

    inf = pd.DataFrame(rows)
    inf["rel_cost_vs_I00"] = inf["p50_ms"] / base_lat["p50"]
    inf["cost_class"] = inf["exp_id"].map(_cost_class)
    inf["realtime_ok"] = inf["p95_ms"] <= LATENCY_BUDGET_MS
    lat_df = pd.DataFrame(lat_rows)
    cache_csv.parent.mkdir(parents=True, exist_ok=True)
    inf.to_csv(cache_csv, index=False)
    lat_df.to_csv(lat_csv, index=False)
    del model, fused
    _empty_cache()
    return inf, lat_df


def inference_tradeoff_summary(inf_df: pd.DataFrame) -> str:
    """Số liệu để trả lời: TTA/ensemble hợp ngoại tuyến, robot nên dùng thứ không tốn thêm chi phí suy luận?"""
    d = inf_df[inf_df.val_macro_f1.notna()]
    i00 = d[d.exp_id == "I00"].iloc[0]
    free = d[d.cost_class == "không tốn thêm"]
    paid = d[d.cost_class.str.startswith("tốn thêm")]
    parts = [f"I00: macro-F1 {i00.val_macro_f1:.4f}, p95 {i00.p95_ms:.1f} ms."]
    if len(free):
        b = free.loc[free.val_macro_f1.idxmax()]
        parts.append(f"Tốt nhất nhóm không tốn thêm: {b.exp_id} ({b.method}) {b.val_macro_f1:.4f} "
                     f"(Δ {b.val_macro_f1 - i00.val_macro_f1:+.4f}), p95 {b.p95_ms:.1f} ms.")
    if len(paid):
        b = paid.loc[paid.val_macro_f1.idxmax()]
        parts.append(f"Tốt nhất nhóm tốn thêm: {b.exp_id} ({b.method}) {b.val_macro_f1:.4f} "
                     f"(Δ {b.val_macro_f1 - i00.val_macro_f1:+.4f}), p95 {b.p95_ms:.1f} ms = x{b.rel_cost_vs_I00:.1f} I00.")
    ok = d[d.realtime_ok]
    parts.append(f"{len(ok)}/{len(d)} phương pháp có p95 ≤ {LATENCY_BUDGET_MS:.0f} ms ở batch 1.")
    return " ".join(parts)


def select_inference(inf_df: pd.DataFrame, override: str | None = None) -> tuple[str, str]:
    """Quy tắc: trong các phương pháp MỘT mô hình (I00-I04), lấy macro-F1 val cao nhất; chỉ dùng nếu hơn
    I00 ít nhất INFER_MIN_GAIN, ngược lại giữ I00. Ensemble không chọn vì chung kết phải train lại ≥ 3 seed
    cho mọi thành viên (vượt ngân sách GPU). Temperature scaling luôn áp dụng thêm (không đổi accuracy)."""
    if override:
        return override, f"chọn thủ công: {override}"
    single = inf_df[inf_df.exp_id.isin(METHODS.keys()) & inf_df.val_macro_f1.notna()]
    i00 = single.loc[single.exp_id == "I00", "val_macro_f1"].item()
    best = single.loc[single.val_macro_f1.idxmax()]
    if best.val_macro_f1 - i00 >= INFER_MIN_GAIN:
        return best.exp_id, (f"{best.exp_id} ({best.method}) macro-F1 val {best.val_macro_f1:.4f} > I00 {i00:.4f} "
                             f"(Δ={best.val_macro_f1 - i00:+.4f} ≥ {INFER_MIN_GAIN}); chi phí x{best.rel_cost_vs_I00:.1f}.")
    return "I00", (f"Không phương pháp một mô hình nào hơn I00 ({i00:.4f}) từ {INFER_MIN_GAIN} trở lên "
                   f"(tốt nhất {best.exp_id} {best.val_macro_f1:.4f}); giữ 1 view.")


# ============================================================================= Bước 4
def final_predict(summary: dict, method_code: str, exp_id: str, pred_dir: str | Path, device,
                  calibrate: bool = True, realtime_id: str | None = None) -> dict:
    """Chạy VAL và TEST đúng một lần cho một seed, ghi file dự đoán đúng định dạng eval.py.

    - <exp_id>_seed<k>_val.csv / _test.csv : phương pháp `method_code` (+ temperature scaling nếu calibrate,
      T khớp trên VAL của chính seed đó)
    - <exp_id>uncal_seed<k>_test.csv        : cùng phương pháp, chưa temperature scaling (I4a)
    - <realtime_id>_seed<k>_test.csv        : 1 view + TS (cấu hình thời gian thực, I5)
    Nếu file test đã tồn tại thì KHÔNG chạy lại (quy tắc: test một lần mỗi seed).
    """
    pred_dir = Path(pred_dir)
    seed = summary["seed"]
    out_test = pred_dir / f"{exp_id}_seed{seed}_test.csv"
    info_path = Path(summary["run_dir"]) / f"final_{exp_id}.json"
    if out_test.exists():
        print(f"{out_test.name} đã có -> không chạy test lần nữa")
        return _read_json(info_path) if info_path.exists() else {}

    cfg = load_config(summary)
    model = load_trained(summary, device)
    _, val_df, test_df = D.load_split(cfg.labels_dir, cfg.fold)
    meth = METHODS[method_code]
    info = {"exp_id": exp_id, "seed": seed, "source_run": f"{summary['exp_id']} seed{seed}",
            "method": method_code, "calibrate": calibrate}

    # --- VAL: khớp T (chỉ dùng val)
    nv, yv, views_v = method_logits(model, cfg, val_df, meth, device)
    pv, zv = combine(views_v, meth)
    T = I.fit_temperature(zv, yv) if calibrate else 1.0
    save_predictions(pred_dir / f"{exp_id}_seed{seed}_val.csv", nv, yv, I.apply_temperature(zv, T) if calibrate else pv)
    info["T"] = T
    if realtime_id:
        nv1, yv1, v1 = method_logits(model, cfg, val_df, METHODS["I00"], device)
        T1 = I.fit_temperature(v1[0], yv1)
        info["T_realtime"] = T1

    # --- TEST: đúng một lần
    nt, yt, views_t = method_logits(model, cfg, test_df, meth, device)
    pt, zt = combine(views_t, meth)
    if calibrate:
        save_predictions(pred_dir / f"{exp_id}uncal_seed{seed}_test.csv", nt, yt, pt)
        save_predictions(out_test, nt, yt, I.apply_temperature(zt, T))
    else:
        save_predictions(out_test, nt, yt, pt)
    if realtime_id:
        if method_code == "I00":
            zt1 = views_t[0]
        else:
            _, _, v1t = method_logits(model, cfg, test_df, METHODS["I00"], device)
            zt1 = v1t[0]
        save_predictions(pred_dir / f"{realtime_id}_seed{seed}_test.csv", nt, yt, I.apply_temperature(zt1, T1))
    _write_json(info_path, info)
    del model
    _empty_cache()
    print(f"[{exp_id} seed{seed}] đã ghi dự đoán val/test ({method_code}, T={T:.3f})")
    return info


def step4_final(paths: dict, backbone: str, recipe: dict, method_code: str, pred_dir: str | Path, device,
                seeds=(0, 1, 2), **overrides) -> dict:
    """Train (hoặc dùng lại) F01 và mốc T00 với các seed, rồi ghi dự đoán test mỗi seed một lần."""
    out = {"F01": [], "T00": []}
    for seed in seeds:
        fs = run_or_load(Config(exp_id="F01", backbone=backbone, desc="final", seed=seed,
                                **{**overrides, **recipe}, **paths))
        bs = _t00(paths, backbone, seed, **overrides)
        out["F01"].append(fs)
        out["T00"].append(bs)
    for fs, bs in zip(out["F01"], out["T00"]):
        final_predict(fs, method_code, "F01", pred_dir, device, calibrate=True, realtime_id="F01rt")
        final_predict(bs, "I00", "T00", pred_dir, device, calibrate=False)
    return out


def final_configs(backbone: str, recipe_id: str, recipe: dict, method: str) -> dict:
    """Mô tả cấu hình (backbone + công thức + suy luận) của từng nhóm trong sheet Final."""
    rec = f"{recipe_id} {recipe}" if recipe else "T00 (công thức nền)"
    meth = METHODS[method]["label"]
    return {"F01": f"{backbone} + {rec} + {method} ({meth}) + temperature scaling",
            "F01uncal": f"{backbone} + {rec} + {method} ({meth}), chưa temperature scaling",
            "F01rt": f"{backbone} + {rec} + I00 (1 view) + temperature scaling [thời gian thực]",
            "T00": f"{backbone} + T00 (công thức nền) + I00 (1 view) [mốc]"}


def final_table(pred_dir: str | Path, groups=("F01", "T00", "F01rt", "F01uncal"),
                configs: dict | None = None) -> pd.DataFrame:
    """Bảng Final: mỗi seed một dòng + dòng mean ± std, tính từ file dự đoán bằng eval.compute_metrics.
    configs: {nhóm: mô tả cấu hình} (xem final_configs)."""
    configs = configs or {}
    import eval as ev
    pred_dir = Path(pred_dir)
    rows = []
    for g in groups:
        files = sorted(pred_dir.glob(f"{g}_seed*_test.csv"))
        if not files:
            continue
        per = []
        for f in files:
            p = ev.read_pred(str(f))
            mt = compute_metrics(p.y_true, p.y_pred, p.probs)
            vf = pred_dir / f.name.replace("_test.csv", "_val.csv")
            mv = None
            if vf.exists():
                pv = ev.read_pred(str(vf))
                mv = compute_metrics(pv.y_true, pv.y_pred, pv.probs)
            r = {"exp_id": g, "config": configs.get(g, ""), "seed": p.seed, "val_macro_f1": mv["macro_f1"] if mv else np.nan,
                 "test_macro_f1": mt["macro_f1"], "test_top1": mt["top1"], "test_balanced_acc": mt["balanced_acc"],
                 "test_ece": mt["ece"], **{f"test_recall_{k.split()[0].lower()}": mt["recall"][i] for k, i in RARE.items()}}
            per.append(r)
        rows += per
        agg = {"exp_id": g, "config": configs.get(g, ""), "seed": f"mean ± std ({len(per)} seed)"}
        for k in per[0]:
            if k in ("exp_id", "config", "seed"):
                continue
            vals = [r[k] for r in per]
            if all(pd.isna(v) for v in vals):  # vd F01rt/F01uncal không có file val
                agg[k] = ""
                continue
            mu, sd = mean_std(vals)
            agg[k] = f"{mu:.4f} ± {sd:.4f}"
        rows.append(agg)
    return pd.DataFrame(rows)


# ============================================================================= Bước 5: hình và xlsx
def plot_backbones(bb_df: pd.DataFrame, path: str | Path) -> None:
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(1, 2, figsize=(11, 4))
    for _, r in bb_df.iterrows():
        ax[0].scatter(r.latency_b1_p50_ms, r.val_macro_f1, s=30 + 8 * r.params_M)
        ax[0].annotate(r.backbone, (r.latency_b1_p50_ms, r.val_macro_f1), fontsize=8, xytext=(4, 4),
                       textcoords="offset points")
        ax[1].scatter(r.gmacs, r.latency_b1_p50_ms)
        ax[1].annotate(r.backbone, (r.gmacs, r.latency_b1_p50_ms), fontsize=8, xytext=(4, 4), textcoords="offset points")
    ax[0].set(xlabel="độ trễ p50 batch 1 (ms)", ylabel="macro-F1 val", title="Backbone: chất lượng vs độ trễ (cỡ = #params)")
    ax[1].set(xlabel="GMAC", ylabel="độ trễ p50 batch 1 (ms)", title="FLOPs có dự đoán được độ trễ?")
    for a in ax:
        a.grid(alpha=0.3)
    fig.tight_layout()
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=130)
    plt.close(fig)


def plot_tradeoff(inf_df: pd.DataFrame, path: str | Path) -> None:
    import matplotlib.pyplot as plt
    d = inf_df[inf_df.val_macro_f1.notna() & inf_df.p50_ms.notna()]
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.scatter(d.p95_ms, d.val_macro_f1)
    for _, r in d.iterrows():
        ax.annotate(r.exp_id, (r.p95_ms, r.val_macro_f1), fontsize=8, xytext=(4, 3), textcoords="offset points")
    ax.axvline(LATENCY_BUDGET_MS, color="red", ls="--", lw=1, label=f"ngân sách {LATENCY_BUDGET_MS:.0f} ms")
    ax.set_xscale("log")
    ax.set(xlabel="độ trễ p95 batch 1 (ms, thang log)", ylabel="macro-F1 val",
           title="Đánh đổi độ chính xác – độ trễ của các phương pháp suy luận")
    ax.legend()
    ax.grid(alpha=0.3, which="both")
    fig.tight_layout()
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=130)
    plt.close(fig)


def plot_confusion(pred_glob_dir: str | Path, group: str, path: str | Path) -> np.ndarray:
    """Ma trận nhầm lẫn TỔNG qua các seed (số ảnh), hàng = nhãn thật, cột = dự đoán."""
    import matplotlib.pyplot as plt
    import eval as ev
    cms = []
    for f in sorted(Path(pred_glob_dir).glob(f"{group}_seed*_test.csv")):
        p = ev.read_pred(str(f))
        cms.append(ev.confusion_matrix(p.y_true, p.y_pred))
    cm = np.sum(cms, 0)
    fig, ax = plt.subplots(figsize=(8, 7))
    norm = cm / cm.sum(1, keepdims=True)
    ax.imshow(norm, cmap="Blues")
    for i in range(cm.shape[0]):
        for j in range(cm.shape[1]):
            ax.text(j, i, cm[i, j], ha="center", va="center", fontsize=8,
                    color="white" if norm[i, j] > 0.5 else "black")
    ax.set_xticks(range(len(D.CLASS_NAMES)), D.CLASS_NAMES, rotation=40, ha="right", fontsize=8)
    ax.set_yticks(range(len(D.CLASS_NAMES)), D.CLASS_NAMES, fontsize=8)
    ax.set(xlabel="dự đoán", ylabel="nhãn thật", title=f"Ma trận nhầm lẫn test – {group} (tổng {len(cms)} seed)")
    fig.tight_layout()
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=130)
    plt.close(fig)
    return cm


def plot_errors(pred_file: str | Path, images_dir: str | Path, path: str | Path,
                pairs=((0, 7), (7, 0)), n: int = 6) -> None:
    """Ảnh test bị đoán sai cho các cặp (thật -> dự đoán), mặc định Chinee Apple <-> Snake Weed."""
    import matplotlib.pyplot as plt
    from PIL import Image
    import eval as ev
    p = ev.read_pred(str(pred_file))
    fig, axes = plt.subplots(len(pairs), n, figsize=(2.2 * n, 2.6 * len(pairs)), squeeze=False)
    for r, (t, q) in enumerate(pairs):
        idx = np.where((p.y_true == t) & (p.y_pred == q))[0][:n]
        for c in range(n):
            ax = axes[r, c]
            ax.axis("off")
            if c < len(idx):
                i = idx[c]
                ax.imshow(Image.open(Path(images_dir) / p.filenames[i]))
                ax.set_title(f"conf {p.probs[i].max():.2f}", fontsize=7)
        axes[r, 0].text(-0.1, 0.5, f"{D.CLASS_NAMES[t]}\n→ {D.CLASS_NAMES[q]}\n({len(np.where((p.y_true == t) & (p.y_pred == q))[0])} ảnh)",
                        transform=axes[r, 0].transAxes, ha="right", va="center", fontsize=8)
    fig.tight_layout()
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=130, bbox_inches="tight")
    plt.close(fig)


def top_confusions(cm: np.ndarray, k: int = 5) -> pd.DataFrame:
    """k cặp (nhãn thật -> dự đoán) bị nhầm nhiều nhất từ ma trận nhầm lẫn (số ảnh, tổng qua seed)."""
    off = cm.astype(float).copy()
    np.fill_diagonal(off, -1)
    rows = []
    for idx in np.argsort(off, axis=None)[::-1][:k]:
        t, q = np.unravel_index(idx, cm.shape)
        if off[t, q] <= 0:
            break
        rows.append({"true": D.CLASS_NAMES[t], "pred": D.CLASS_NAMES[q], "n_images": int(cm[t, q]),
                     "pct_of_true_class": float(cm[t, q] / cm[t].sum()), "true_idx": int(t), "pred_idx": int(q)})
    return pd.DataFrame(rows)


def per_class_table(eval_out: str | Path, tags=("F01", "T00", "F01rt")) -> pd.DataFrame:
    """Sheet PerClass từ file <tag>_per_class.csv do `eval.py score --out` ghi ra (số ảnh test = support)."""
    frames = []
    for t in tags:
        f = Path(eval_out) / f"{t}_per_class.csv"
        if f.exists():
            d = pd.read_csv(f)
            d.insert(0, "config", t)
            frames.append(d)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def hard_class_table(pc_df: pd.DataFrame) -> pd.DataFrame:
    """Precision/recall/F1 (mean ± std qua seed) của Chinee Apple và Snake Weed, lấy từ output eval.py."""
    d = pc_df[_is_rare(pc_df["class"])].copy()
    for k in ("precision", "recall", "f1"):
        d[k] = [f"{m:.4f} ± {sd:.4f}" for m, sd in zip(d[f"{k}_mean"], d[f"{k}_std"])]
    paper = {"chinee apple": 0.885, "snake weed": 0.888}
    d["paper_recall"] = d["class"].astype(str).str.lower().map(paper)
    return d[["config", "class", "support", "precision", "recall", "f1", "paper_recall"]]


def summary_table(bb_df, tr_df, inf_df, final_df) -> pd.DataFrame:
    """Top 10 cấu hình theo macro-F1 val kèm chi phí/độ trễ, rồi dòng chung kết và mốc trên test."""
    lat_of_bb = dict(zip(bb_df.backbone, bb_df.latency_b1_p95_ms))
    rows = []
    for r in bb_df.itertuples():
        rows.append({"exp_id": r.exp_id, "loại": "backbone", "mô tả": r.backbone, "val_macro_f1": r.val_macro_f1,
                     "val_top1": r.val_top1, "p95_b1_ms": r.latency_b1_p95_ms, "rel_cost": np.nan})
    for r in tr_df.itertuples():
        if r.exp_id == "T00":
            continue
        rows.append({"exp_id": r.exp_id, "loại": "huấn luyện", "mô tả": r.change_vs_T00,
                     "val_macro_f1": r.val_macro_f1, "val_top1": r.val_top1,
                     "p95_b1_ms": lat_of_bb.get(r.backbone, np.nan), "rel_cost": 1.0})
    for r in inf_df[inf_df.val_macro_f1.notna()].itertuples():
        rows.append({"exp_id": r.exp_id, "loại": "suy luận", "mô tả": r.method, "val_macro_f1": r.val_macro_f1,
                     "val_top1": r.val_top1, "p95_b1_ms": r.p95_ms, "rel_cost": r.rel_cost_vs_I00})
    top = pd.DataFrame(rows).sort_values("val_macro_f1", ascending=False).head(10)
    fin = final_df[final_df.seed.astype(str).str.startswith("mean")][
        ["exp_id", "config", "seed", "val_macro_f1", "test_macro_f1", "test_top1", "test_ece"]]
    fin = fin.rename(columns={"config": "mô tả", "seed": "số seed"}).assign(loại="chung kết (test)")
    return pd.concat([top, fin], ignore_index=True)


def _rows_to_highlight(df: pd.DataFrame, rule) -> list[int]:
    """rule: tên cột (lấy max) | ("min", cột) | callable(df) -> list chỉ số dòng."""
    if callable(rule):
        return list(rule(df))
    how, col = ("max", rule) if isinstance(rule, str) else rule
    vals = pd.to_numeric(df[col], errors="coerce") if col in df.columns else pd.Series(dtype=float)
    if not vals.notna().any():
        return []
    return [int(vals.idxmax() if how == "max" else vals.idxmin())]


DEFAULT_HIGHLIGHT = {
    "Backbones": "val_macro_f1",
    "Training": "val_macro_f1",
    "Inference": "val_macro_f1",
    "Final": lambda d: d.index[(d.exp_id == "F01") & d.seed.astype(str).str.startswith("mean")],
    "PerClass": lambda d: d.index[(d["config"] == "F01") & _is_rare(d["class"])],
    "Latency": lambda d: d.index[d.batch == 1][pd.to_numeric(d.loc[d.batch == 1, "p95_ms"]).argmin():][:1],
    "Summary": lambda d: ([pd.to_numeric(d.val_macro_f1, errors="coerce").idxmax()]
                          + list(d.index[d.exp_id.isin(["F01", "T00"]) & (d["loại"] == "chung kết (test)")])),
}


def write_results_xlsx(path: str | Path, sheets: dict[str, pd.DataFrame], highlight: dict | None = None) -> None:
    """Ghi results.xlsx: cố định hàng tiêu đề, 4 chữ số thập phân, tự giãn cột, tô nổi bật dòng tốt nhất của
    mỗi sheet (highlight mặc định DEFAULT_HIGHLIGHT; giá trị: cột lấy max, ("min", cột) hoặc hàm chọn dòng)."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    highlight = {**DEFAULT_HIGHLIGHT, **(highlight or {})}
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df = df.reset_index(drop=True)
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for j, col in enumerate(df.columns, 1):
                width = max([len(str(col))] + [len(f"{v:.4f}" if isinstance(v, float) else str(v)) for v in df[col]])
                ws.column_dimensions[get_column_letter(j)].width = min(60, width + 2)
                if pd.api.types.is_float_dtype(df[col]):
                    for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                        row[0].number_format = "0.0000"
            rule = highlight.get(name)
            try:
                idx = _rows_to_highlight(df, rule) if rule is not None and len(df) else []
            except Exception as e:  # tô màu chỉ là trình bày: không được làm hỏng việc ghi file
                print(f"[{name}] bỏ qua tô màu: {e}")
                idx = []
            for i in idx:
                for cell in ws[int(i) + 2]:
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
                    cell.font = Font(bold=True)


def consistency_check(fin_df: pd.DataFrame, pc_df: pd.DataFrame, eval_out: str | Path, tol: float = 1e-6) -> list[str]:
    """So số trong sheet Final/PerClass với JSON do `eval.py score` ghi ra. Trả về danh sách chỗ lệch (rỗng = khớp)."""
    problems = []
    for tag in fin_df.exp_id.unique():
        f = Path(eval_out) / f"{tag}_summary.json"
        if not f.exists():
            problems.append(f"thiếu {f.name} (chưa chạy eval.py score cho {tag})")
            continue
        ev_sum = _read_json(f)
        per = fin_df[(fin_df.exp_id == tag) & ~fin_df.seed.astype(str).str.startswith("mean")]
        for mine, theirs in (("test_macro_f1", "macro_f1"), ("test_top1", "top1"), ("test_ece", "ece")):
            a, b = float(pd.to_numeric(per[mine]).mean()), float(ev_sum[theirs]["mean"])
            if abs(a - b) > tol:
                problems.append(f"{tag} {mine}: Final {a:.6f} != eval.py {b:.6f}")
        if sorted(int(x) for x in per.seed) != sorted(ev_sum["seeds"]):
            problems.append(f"{tag}: seed trong Final {sorted(per.seed)} != eval.py {ev_sum['seeds']}")
        for name, i in RARE.items():
            sub = (pc_df[(pc_df["config"] == tag) & (pc_df["class"].astype(str).str.lower() == name.lower())]
                   if len(pc_df) else pc_df)
            if len(sub):
                a, b = float(sub["recall_mean"].iloc[0]), float(ev_sum["recall"]["mean"][i])
                if abs(a - b) > tol:
                    problems.append(f"{tag} recall {name}: PerClass {a:.6f} != eval.py {b:.6f}")
    return problems


def check_curves(out_dir: str | Path, curves_dir: str | Path) -> list[str]:
    """Mỗi lần train (B, T, F; mọi seed) phải có ảnh curves/<exp_id>_<mota>[_seedk].png. Trả về danh sách thiếu."""
    missing = []
    for cj in sorted(Path(out_dir).glob("*/seed*/config.json")):
        if not (cj.parent / "summary.json").exists():
            continue
        name = curve_path(Config(**_read_json(cj))).name
        if not (Path(curves_dir) / name).exists():
            missing.append(name)
    return missing


def contributions(bb_df, tr_df, inf_df, backbone: str, recipe_id: str, method: str, fin_df, noise: dict | None) -> pd.DataFrame:
    """Bảng đóng góp: backbone / công thức / suy luận (trên val, 1 seed) và chung kết vs mốc (test, nhiều seed)."""
    thr = (noise or {}).get("threshold", NOISE_F1)
    ref = bb_df[bb_df.exp_id == "B01"].iloc[0]
    chosen = bb_df[bb_df.backbone == backbone].iloc[0]
    t00 = tr_df[(tr_df.exp_id == "T00") & (tr_df.seed == 0)].iloc[0]
    rec = tr_df[(tr_df.exp_id == recipe_id) & (tr_df.seed == 0)].iloc[0]
    i00 = inf_df[inf_df.exp_id == "I00"].iloc[0]
    meth = inf_df[inf_df.exp_id == method].iloc[0]
    rows = [
        {"yếu tố": "backbone", "so sánh": f"{chosen.backbone} vs {ref.backbone} (B01, mốc)", "tập": "val, 1 seed",
         "delta_macro_f1": chosen.val_macro_f1 - ref.val_macro_f1},
        {"yếu tố": "công thức huấn luyện", "so sánh": f"{recipe_id} vs T00", "tập": "val, 1 seed",
         "delta_macro_f1": rec.val_macro_f1 - t00.val_macro_f1},
        {"yếu tố": "suy luận", "so sánh": f"{method} vs I00", "tập": "val, 1 seed",
         "delta_macro_f1": meth.val_macro_f1 - i00.val_macro_f1},
    ]
    for r in rows:
        r["noise_threshold"] = thr
        r["kết luận"] = _vs_noise(r["delta_macro_f1"], thr)
    fm = fin_df[fin_df.seed.astype(str).str.startswith("mean")].set_index("exp_id")
    if {"F01", "T00"} <= set(fm.index):
        per = fin_df[~fin_df.seed.astype(str).str.startswith("mean")]
        f = pd.to_numeric(per[per.exp_id == "F01"].test_macro_f1)
        b = pd.to_numeric(per[per.exp_id == "T00"].test_macro_f1)
        d = f.mean() - b.mean()
        sd = max(f.std(ddof=1), b.std(ddof=1))
        rows.append({"yếu tố": "tổng (chung kết vs mốc)", "so sánh": "F01 vs T00+I00", "tập": f"test, {len(f)} seed",
                     "delta_macro_f1": d, "noise_threshold": sd,
                     "kết luận": ("vượt nhiễu (Δ > std lớn hơn của hai nhóm)" if d > sd else
                                  "không phân biệt được (Δ ≤ std)" if d > -sd else "kém hơn rõ")})
    return pd.DataFrame(rows)


def experiment_index(out_dir: str | Path) -> pd.DataFrame:
    """Phụ lục: mọi lần train, điểm khác so với T00 (seed 0), lần dùng lại và đường dẫn config.json."""
    out_dir = Path(out_dir)
    base_p = out_dir / "T00" / "seed0" / "config.json"
    base = _recipe(_read_json(base_p)) if base_p.exists() else {}
    rows = []
    for cj in sorted(out_dir.glob("*/seed*/config.json")):
        sp = cj.parent / "summary.json"
        if not sp.exists():
            continue
        c, sm = _read_json(cj), _read_json(sp)
        diff = {k: v for k, v in _recipe(c).items() if k not in ("seed",) and base.get(k) != v}
        rows.append({"exp_id": c["exp_id"], "seed": c["seed"], "backbone": c["backbone"],
                     "khác T00": ", ".join(f"{k}={v}" for k, v in diff.items()) or "-",
                     "val_macro_f1": sm["val_macro_f1"], "best_epoch": sm["best_epoch"],
                     "reused_from": sm.get("reused_from", ""), "config": str(cj.relative_to(out_dir))})
    return pd.DataFrame(rows)


# ============================================================================= tiện ích cho notebook
def run_summary(paths: dict, exp_id: str, seed: int = 0) -> dict:
    return _read_json(Path(paths["out_dir"]) / exp_id / f"seed{seed}" / "summary.json")


def add_combo_row(tr_df: pd.DataFrame, combo_summary: dict, reason: str, noise: dict | None = None) -> pd.DataFrame:
    base = tr_df.loc[(tr_df.exp_id == "T00") & (tr_df.seed == 0), "val_macro_f1"].item()
    row = _training_row(combo_summary, COMBO_ID, "kết hợp", reason, base)
    df = pd.concat([tr_df[tr_df.exp_id != COMBO_ID], pd.DataFrame([row])], ignore_index=True)
    return annotate_noise(df, noise)


def ensemble_candidates(paths: dict, *dfs: pd.DataFrame, k: int = 3) -> list[dict]:
    """k lần chạy có macro-F1 val cao nhất (Bước 1-2), bỏ các lần 'dùng lại' để không trùng mô hình."""
    allr = pd.concat([d[["exp_id", "val_macro_f1"]] for d in dfs]).drop_duplicates("exp_id")
    out = []
    for e in allr.sort_values("val_macro_f1", ascending=False).exp_id:
        s = run_summary(paths, e)
        if not s.get("reused_from"):
            out.append(s)
        if len(out) == k:
            break
    return out


def realtime_p95(inf_df: pd.DataFrame) -> tuple[float, str]:
    """p95 batch 1 của cấu hình thời gian thực F01rt (1 view). Dự đoán chung kết chạy với AMP trên GPU,
    nên lấy dòng I08_amp nếu có (cùng dtype với file dự đoán), ngược lại lấy I00 (FP32)."""
    for code in ("I08_amp", "I00"):
        r = inf_df[inf_df.exp_id == code]
        if len(r) and pd.notna(r.p95_ms.iloc[0]):
            return float(r.p95_ms.iloc[0]), code
    raise ValueError("chưa có số đo độ trễ I00")


def env_info() -> dict:
    import platform

    import timm
    import torchvision
    return {"python": platform.python_version(), "torch": torch.__version__, "torchvision": torchvision.__version__,
            "timm": timm.__version__, "numpy": np.__version__, "pandas": pd.__version__,
            "gpu": torch.cuda.get_device_name(0) if torch.cuda.is_available() else "CPU",
            "cuda": torch.version.cuda}


def _fmt(v, digits: int = 4) -> str:
    if v is None:
        return ""
    if isinstance(v, (float, np.floating)):
        return "" if np.isnan(v) else f"{v:.{digits}f}"
    return str(v).replace("|", "/")


def md_table(df: pd.DataFrame, cols: list[str] | None = None, digits: int = 4) -> str:
    """Bảng Markdown (không cần thư viện tabulate)."""
    d = df[[c for c in cols if c in df.columns]] if cols else df
    lines = ["| " + " | ".join(map(str, d.columns)) + " |", "|" + "---|" * len(d.columns)]
    lines += ["| " + " | ".join(_fmt(v, digits) for v in row) + " |" for row in d.itertuples(index=False)]
    return "\n".join(lines)


REPORT_TEMPLATE = """# Báo cáo Lab Day 2 — DeepWeeds · Nguyễn Văn Quốc Việt · 2A202602973

> Bản nháp sinh tự động từ kết quả chạy thật: mọi bảng lấy từ `results.xlsx` / output của `eval.py`, mọi lựa chọn
> kèm lý do theo quy tắc trên val. Các mục **TODO** là phần nhận xét, phân tích bạn tự viết.
> Kiểm tra khớp số (Final/PerClass so với `eval.py score`): {consistency}

## 1. Tóm tắt
- Cấu hình tốt nhất F01: backbone `{backbone}` + công thức `{recipe_id}` + suy luận `{method}` + temperature scaling.
- Test ({n_seed}): macro-F1 **{f_mf1}**, top-1 **{f_top1}**, ECE {f_ece}.
- Mốc T00 + I00: macro-F1 {b_mf1}, top-1 {b_top1}.
- TODO: 3–5 dòng kết luận chính (dựa vào bảng đóng góp ở mục 7).

## 2. Dữ liệu và thiết lập
- DeepWeeds, fold 0 chia sẵn (train_subset0 / val_subset0 / test_subset0), không sửa, không lọc, không chia lại.
{split_md}- Chỉ số chính: macro-F1 (9 lớp, `eval.compute_metrics`); phụ: top-1, balanced accuracy, F1/recall từng lớp, ECE 15 bin.
- Công thức nền T00: ImageNet-1k pretrained, tinh chỉnh toàn bộ, RandomResizedCrop 224 + lật ngang, AdamW (LR backbone 1e-4,
  head 1e-3, weight decay 0,05 trừ norm/bias), warmup 1 epoch + cosine theo bước, CE, batch 64, {epochs} epoch, AMP,
  channels_last, chọn checkpoint theo macro-F1 val (hòa lấy epoch sớm hơn).
- Val/test: resize 256 → center crop 224, chuẩn hoá theo mean/std của từng bộ trọng số. Seed 0 cho quét sàng; 0, 1, 2 cho chung kết.
- Môi trường: {env}.
{pipeline_md}
![Phân bố lớp](figures/eda_class_distribution.png)

![Ảnh mẫu](figures/eda_samples.png)

TODO: nhận xét mất cân bằng (Negatives ≈ 52%) và cặp loài dễ nhầm bằng mắt thường; Negatives trông ra sao.

## 3. So sánh backbone (Bước 1, 1 seed)
{bb_table}

Phân tích tự động từ `history.csv` (epochs_to_99pct: epoch đầu đạt 99% macro-F1 tốt nhất; overfit: val loss
tăng > 0,05 sau điểm thấp nhất trong khi train loss vẫn giảm):

{bb_analysis}

Tương quan thứ hạng (5 backbone, chỉ tham khảo): {bb_corr}

![backbones](figures/backbones.png)

![backbone curves](figures/backbones_curves.png)

**Lựa chọn (quy tắc trên val):** {why_bb}

TODO: backbone nào hội tụ nhanh nhất / có overfit không; thứ hạng trên DeepWeeds so với ImageNet; FLOPs có dự đoán
được thời gian train và độ trễ không (slide trang 43). Đây là kết quả 1 seed: chênh lệch < {noise_b1} là không phân biệt được.

## 4. Công thức huấn luyện (Bước 2, 1 seed mỗi biến thể)
Nhiễu đo được: T00 với seed {noise_seeds} có macro-F1 val {noise_vals}, mean {noise_mean:.4f}, std (ddof=1) {noise_std:.4f}.
Ngưỡng dùng để kết luận = max(std, {noise_floor}) = **{noise_thr:.4f}**: |Δ| ≤ ngưỡng thì ghi "không phân biệt được".

{tr_table}

Mỗi T0x khác T00 đúng một yếu tố. {why_combo}

**Công thức chung kết:** {why_recipe}

TODO: yếu tố nào giúp / không giúp và vì sao (liên hệ slide); sampler (T09) khác loss có trọng số (T07) thế nào;
hiệu ứng có cộng dồn khi kết hợp (T10) không.

## 5. Suy luận (Bước 3, trên val)
{inf_table}

![tradeoff](figures/tradeoff.png)

Độ trễ: warmup 10 lần, `torch.cuda.synchronize()` trước và sau, 100 lần đo, báo p50/p95/p99, chỉ đo forward
(không tính tiền xử lý), GPU {gpu}, torch {torch}. Thông lượng đo ở batch 32. Bảng đầy đủ (batch 1 và 32,
FP32/AMP/FP16, có/không gộp BN) ở sheet Latency.

Temperature scaling (I07): ECE val trước {ece_before}, sau {ece_after} (in-sample), {ece_cf} (cross-fit hai nửa val); T = {temp}.

**Đánh đổi (số liệu cho câu hỏi "TTA/ensemble hợp ngoại tuyến, robot dùng thứ không tốn thêm"):** {tradeoff}

**Phương pháp cho chung kết:** {why_inf} Temperature scaling (T khớp trên val của từng seed) luôn áp dụng thêm.

TODO: TTA tăng bao nhiêu và tốn bao nhiêu lần độ trễ; EMA vs không EMA (I06 vs I06_raw, cùng lần chạy); soup vs
ensemble seed; độ phân giải kiểm tra (FixRes); dữ liệu có ủng hộ nhận định của slide không.

## 6. Cấu hình tốt nhất và kết quả test (Bước 4)
Tái lập: {final_desc}. Huấn luyện bằng `train.run(Config(...))` với các tham số ở phụ lục; test chạy đúng một lần cho
mỗi seed, sau khi đã chốt mọi lựa chọn trên val. F01rt = cùng mô hình F01, 1 view (thời gian thực, p95 batch 1 =
{rt_p95:.1f} ms theo {rt_src}). F01uncal = F01 chưa temperature scaling.

{fin_table}

Hai lớp khó (test, mean ± std qua seed; recall bài báo chỉ để tham chiếu):

{hard_table}

Các cặp bị nhầm nhiều nhất (F01, tổng các seed):

{confusions}

![confusion](figures/confusion_F01.png)

![errors Chinee Apple ↔ Snake Weed](figures/errors_chinee_snake.png)

![errors cặp nhầm nhiều nhất](figures/errors_top_pair.png)

Tự chấm phần I (`eval.py grade`, đề xuất):

```
{grade}
```

TODO: giả thuyết cho các cặp nhầm (xem ảnh sai); so với bài báo (95,7% / 95,1%; Chinee 88,5%, Snake 88,8%) nhớ rằng
điều kiện huấn luyện khác (100 epoch, augmentation mạnh).

## 7. Kết luận và khuyến nghị
Bảng đóng góp (ngưỡng nhiễu: val dùng ngưỡng đo ở mục 4; test dùng std lớn hơn của hai nhóm seed, như tiêu chí I2):

{contrib}

- TODO: cấu hình tốt nhất, tốt hơn mốc bao nhiêu, có vượt std không (dòng cuối bảng trên).
- TODO: yếu tố đóng góp nhiều nhất (backbone, huấn luyện hay suy luận).
- TODO: triển khai trên robot với ngân sách 30–100 ms/khung: chọn gì và vì sao (F01rt, sheet Latency).

## 8. Hạn chế và việc tiếp theo
- Quét sàng Bước 1–2 chỉ 1 seed mỗi biến thể; nhiễu đo bằng 3 seed của T00; chung kết 3 seed; chỉ fold 0.
- Chia ngẫu nhiên, không theo địa điểm, nên điểm test có thể lạc quan khi gặp địa điểm/mùa/góc chụp/ánh sáng mới.
- Giảm bớt do ngân sách GPU (một session Kaggle 12 giờ): {epoch_note}5 backbone, ablation trên 1 backbone. T00 seed 0 dùng
  lại lần chạy Bước 1, F01 seed 0 dùng lại lần chạy tốt nhất Bước 2 (cùng cấu hình, cùng seed; cột reused_from).
- TODO: thí nghiệm thất bại hoặc bất thường (nếu có); việc tiếp theo (nhiều fold, chưng cất, thích ứng miền...).

## 9. Phụ lục
Danh sách mọi lần train (khác T00 ở đâu, config đầy đủ tại `runs/<exp_id>/seed<k>/config.json`):

{exp_index}

- Bảng đầy đủ: `results.xlsx`. Notebook: `code/lab_day2.ipynb` (link Kaggle: TODO).
"""


def _get_row(df: pd.DataFrame, code: str, col: str):
    r = df[df.exp_id == code]
    return r[col].iloc[0] if len(r) and col in r else float("nan")


def write_report_draft(path: str | Path, *, env: dict, split_stats: dict | None, backbone: str, bb_df, why_bb,
                       tr_df, why_combo, recipe_id: str, why_recipe, inf_df, method: str, why_inf, fin_df, pc_df,
                       rt_p95: float, rt_src: str, epochs: int, final_desc: str = "", grade_text: str = "",
                       pipeline_checks: dict | None = None, bb_corr: dict | None = None, noise: dict | None = None,
                       tradeoff: str = "", confusions: pd.DataFrame | None = None, contrib: pd.DataFrame | None = None,
                       exp_index: pd.DataFrame | None = None, consistency: list[str] | None = None,
                       overwrite: bool = False) -> Path:
    """Bản nháp report.md theo dàn ý GUIDE mục 6.3: điền sẵn BẢNG SỐ LIỆU thật và lý do chọn cấu hình;
    phần nhận xét để dạng TODO. Không ghi đè file đã có (trừ khi overwrite=True)."""
    path = Path(path)
    if path.exists() and not overwrite:
        print(f"{path} đã có -> không ghi đè (overwrite=True để ghi lại)")
        return path
    fmean = fin_df[fin_df.seed.astype(str).str.startswith("mean")].set_index("exp_id")
    get = lambda g, c: fmean.loc[g, c] if g in fmean.index and c in fmean.columns else "?"  # noqa: E731
    split_md = ""
    if split_stats:
        split_md = (f"- Số ảnh: train {split_stats['n']['train']}, val {split_stats['n']['val']}, "
                    f"test {split_stats['n']['test']} ("
                    + ", ".join(f"{k} {v:.2%}" for k, v in split_stats["frac"].items())
                    + f"). Giao từng cặp: {split_stats['overlap']}; hợp ba tập: {split_stats['union']}; "
                    f"file thiếu: {split_stats['missing']}; nhãn lệch labels.csv gốc: {split_stats.get('label_mismatch')}.\n"
                    f"- Tỉ lệ lớp nhiều nhất / ít nhất: {split_stats.get('imbalance_ratio', float('nan')):.2f}. "
                    "Số ảnh mỗi lớp trong từng tập:\n\n"
                    + md_table(pd.DataFrame(split_stats["per_class"]).rename_axis("lớp").reset_index(), digits=0)
                    + "\n\n")
        for r in split_stats.get("label_mismatch_rows") or []:
            lab = r["label_labels_csv"]
            split_md += (f"- Dữ liệu gốc có nhãn không nhất quán: `{r['Filename']}` ({r['split']}) được ghi lớp "
                         f"{r['label_split_csv']} ({D.CLASS_NAMES[r['label_split_csv']]}) trong file split nhưng "
                         + (f"lớp {lab} ({D.CLASS_NAMES[lab]})" if lab is not None else "không có")
                         + " trong labels.csv. Theo S1 giữ nguyên file split (không sửa); "
                         + ("ảnh thuộc train nên không ảnh hưởng chỉ số val/test." if r["split"] == "train" else
                            "ảnh thuộc tập đánh giá: ghi nhận như một hạn chế.") + "\n")
    pipeline_md = ""
    if pipeline_checks:
        pipeline_md = ("- Kiểm tra pipeline (Bước 0, ResNet-50): "
                       + "; ".join(f"{k} = {v}" for k, v in pipeline_checks.items())
                       + f" (kỳ vọng loss ban đầu ≈ ln 9 = {np.log(9):.3f}).\n")
    f01_pc = pc_df[pc_df["config"] == "F01"] if len(pc_df) else pc_df
    txt = REPORT_TEMPLATE.format(
        consistency=("KHỚP" if consistency == [] else "CHƯA KIỂM TRA" if consistency is None
                     else "LỆCH: " + "; ".join(consistency)),
        backbone=backbone, recipe_id=recipe_id, epochs=epochs, method=method, n_seed=get("F01", "seed"),
        epoch_note=(f"{epochs} epoch cho mọi thí nghiệm (GUIDE gợi ý 10–15), " if epochs < 10 else
                    f"{epochs} epoch (bài báo ~100 epoch nên số tuyệt đối có thể thấp hơn bài báo), "),
        f_mf1=get("F01", "test_macro_f1"), f_top1=get("F01", "test_top1"), f_ece=get("F01", "test_ece"),
        b_mf1=get("T00", "test_macro_f1"), b_top1=get("T00", "test_top1"), split_md=split_md, pipeline_md=pipeline_md,
        env=", ".join(f"{k} {v}" for k, v in env.items()), gpu=env.get("gpu"), torch=env.get("torch"),
        noise_b1=NOISE_F1,
        bb_table=md_table(bb_df, ["exp_id", "backbone", "weights_tag", "params_M", "gmacs", "best_epoch",
                                  "val_macro_f1", "val_top1", "train_time_per_epoch_s", "latency_b1_p50_ms",
                                  "latency_b1_p95_ms"]),
        why_bb=why_bb,
        bb_analysis=md_table(bb_df, ["exp_id", "backbone", "best_epoch", "epochs_to_99pct", "val_loss_min_epoch",
                                     "val_loss_rise", "overfit", "train_val_gap_last"]),
        bb_corr=", ".join(f"{k} = {v:.2f}" for k, v in (bb_corr or {}).items()) or "(chưa tính)",
        noise_seeds=(noise or {}).get("seeds", "?"),
        noise_vals=[round(v, 4) for v in (noise or {}).get("val_macro_f1", [])],
        noise_mean=(noise or {}).get("mean", float("nan")), noise_std=(noise or {}).get("std", float("nan")),
        noise_thr=(noise or {}).get("threshold", NOISE_F1), noise_floor=NOISE_FLOOR,
        tr_table=md_table(tr_df, ["exp_id", "axis", "change_vs_T00", "seed", "val_macro_f1", "val_top1",
                                  "delta_vs_T00", "vs_noise", "f1_chinee", "f1_snake", "note"]),
        why_combo=why_combo, why_recipe=why_recipe,
        inf_table=md_table(inf_df, ["exp_id", "method", "K", "img_size", "val_macro_f1", "val_top1", "val_ece",
                                    "p50_ms", "p95_ms", "p99_ms", "images_per_s_b32", "rel_cost_vs_I00",
                                    "cost_class", "realtime_ok", "note"]),
        ece_before=_fmt(_get_row(inf_df, "I07", "val_ece_before")), ece_after=_fmt(_get_row(inf_df, "I07", "val_ece")),
        ece_cf=_fmt(_get_row(inf_df, "I07", "val_ece_crossfit")), temp=_fmt(_get_row(inf_df, "I07", "temperature"), 3),
        tradeoff=tradeoff or "(chưa tính)", why_inf=why_inf,
        final_desc=final_desc or "(xem sheet Final)", rt_p95=rt_p95, rt_src=rt_src,
        fin_table=md_table(fin_df),
        hard_table=md_table(hard_class_table(pc_df)) if len(pc_df) else "(chưa có)",
        confusions=md_table(confusions, ["true", "pred", "n_images", "pct_of_true_class"])
        if confusions is not None and len(confusions) else "(chưa có)",
        grade=grade_text.strip(),
        contrib=md_table(contrib) if contrib is not None and len(contrib) else "(chưa có)",
        exp_index=md_table(exp_index) if exp_index is not None and len(exp_index) else "(chưa có)")
    path.write_text(txt, encoding="utf-8")
    print("Đã ghi bản nháp", path)
    return path
